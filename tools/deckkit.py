"""
deckkit —— 一个极小的 PPTX 生成 + 预览渲染工具集。

思路：用同一份「幻灯片声明式描述」（元素坐标以英寸为单位）同时产出
  1) .pptx（python-pptx）
  2) .png 预览（Pillow + Noto Sans SC）
因此可以在没有 PowerPoint / LibreOffice 的环境里校验排版是否溢出。

约定：1 英寸 = 72 pt。元素坐标 x/y/w/h 单位为英寸，字号单位为 pt。
"""
from __future__ import annotations

import os
from dataclasses import dataclass, field

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ─────────────────────────────── 主题 ───────────────────────────────
T = {
    "ink":      "0F1B2D",   # 深底 / 主标题
    "ink2":     "23344D",
    "body":     "33425C",
    "muted":    "6B7A90",
    "line":     "D9E0EA",
    "soft":     "F2F5FA",
    "soft2":    "E8EEF8",
    "white":    "FFFFFF",
    "accent":   "2F6BE4",   # 主色（蓝）
    "accent_d": "1E4FB8",
    "accent_l": "E4EDFC",
    "orange":   "DE7A33",
    "orange_l": "FBEBDD",
    "green":    "2E9E6B",
    "green_l":  "E2F3EA",
    "amber":    "D89A28",
    "amber_l":  "FBF0DA",
    "red":      "D14B45",
    "red_l":    "FAE6E5",
    "purple":   "6E56CF",
    "purple_l": "ECE7FA",
}

FONT_CJK = "微软雅黑"      # PowerPoint 中使用的字体
FONT_LAT = "Arial"

SW, SH = 13.3333, 7.5      # 16:9，单位英寸
PAD = 0.62                 # 左右安全边距
CONTENT_TOP = 1.92
CONTENT_BOTTOM = 6.92


def F(size, bold=False, color=None):
    return {"t": "", "size": size, "bold": bold, "color": color or T["body"], "space": 0}


def L(text, size=12, bold=False, color=None, space=0, align="l", bullet=None,
      indent=0.0, line=None):
    return {"t": text, "size": size, "bold": bold, "color": color or T["body"],
            "space": space, "align": align, "bullet": bullet, "indent": indent,
            "line": line}


# ─────────────────────────────── 模型 ───────────────────────────────
@dataclass
class Slide:
    kicker: str | None = None
    title: str | None = None
    sub: str | None = None
    notes: str | None = None
    els: list = field(default_factory=list)
    dark: bool = False


class Deck:
    def __init__(self, footer_left: str):
        self.slides: list[Slide] = []
        self.footer_left = footer_left

    # ---------- 结构 ----------
    def slide(self, kicker=None, title=None, sub=None, notes=None, dark=False):
        s = Slide(kicker=kicker, title=title, sub=sub, notes=notes, dark=dark)
        self.slides.append(s)
        return s

    # ---------- 元素原语 ----------
    def rect(self, s, x, y, w, h, fill=None, line=None, lw=1.0, radius=0.10,
             dash=None, shadow=False):
        s.els.append({"k": "rect", "x": x, "y": y, "w": w, "h": h, "fill": fill,
                      "line": line, "lw": lw, "radius": radius, "dash": dash})
        return s.els[-1]

    def text(self, s, x, y, w, h, lines, valign="t", align="l", pad=0.0):
        s.els.append({"k": "text", "x": x, "y": y, "w": w, "h": h, "lines": lines,
                      "valign": valign, "align": align, "pad": pad})
        return s.els[-1]

    def card(self, s, x, y, w, h, title=None, lines=None, fill=None, accent=None,
             title_size=13.5, body_size=11.2, pad=0.16, radius=0.10,
             title_color=None, icon=None):
        """带左侧色条 + 标题 + 正文的卡片。"""
        self.rect(s, x, y, w, h, fill=fill or T["soft"], line=T["line"], radius=radius)
        if accent:
            self.rect(s, x, y, 0.075, h, fill=accent, radius=0.0)
        tx = x + pad + (0.10 if accent else 0)
        tw = w - 2 * pad - (0.10 if accent else 0)
        cy = y + pad * 0.72
        if title:
            t = title if not icon else f"{icon}  {title}"
            self.text(s, tx, cy, tw, 0.30,
                      [L(t, size=title_size, bold=True,
                         color=title_color or T["ink"])])
            cy += 0.34
        if lines:
            self.text(s, tx, cy, tw, y + h - cy - pad * 0.6, lines)
        return s

    def hline(self, s, x, y, w, color=None, lw=1.0, dash=None):
        s.els.append({"k": "hline", "x": x, "y": y, "w": w, "lw": lw,
                      "color": color or T["line"], "dash": dash})

    def vline(self, s, x, y, h, color=None, lw=1.0, dash=None):
        s.els.append({"k": "vline", "x": x, "y": y, "h": h, "lw": lw,
                      "color": color or T["line"], "dash": dash})

    def chevron(self, s, x, y, w, h, text, fill, tcolor=None, size=11.5, bold=True):
        s.els.append({"k": "chevron", "x": x, "y": y, "w": w, "h": h, "text": text,
                      "fill": fill, "tcolor": tcolor or T["white"], "size": size,
                      "bold": bold})

    def arrow(self, s, x, y, w, h, fill=None):
        s.els.append({"k": "arrow", "x": x, "y": y, "w": w, "h": h,
                      "fill": fill or T["muted"]})

    def table(self, s, x, y, w, rows, col_w, header=True, size=10.2,
              row_h=0.34, head_fill=None, zebra=True):
        s.els.append({"k": "table", "x": x, "y": y, "w": w, "rows": rows,
                      "col_w": col_w, "header": header, "size": size,
                      "row_h": row_h, "head_fill": head_fill or T["ink"],
                      "zebra": zebra})

    def kpi(self, s, x, y, w, h, value, label, color=None, vsize=30):
        self.rect(s, x, y, w, h, fill=T["white"], line=T["line"], radius=0.10)
        self.rect(s, x, y, w, 0.055, fill=color or T["accent"], radius=0.0)
        self.text(s, x, y + 0.16, w, 0.52, [L(value, size=vsize, bold=True,
                                              color=color or T["accent"], align="c")])
        self.text(s, x + 0.08, y + h - 0.46, w - 0.16, 0.40,
                  [L(label, size=9.6, color=T["muted"], align="c")])

    # ---------- 输出 ----------
    def emit(self, path):
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
        blank = prs.slide_layouts[6]
        for i, s in enumerate(self.slides):
            sl = prs.slides.add_slide(blank)
            self._chrome_pptx(sl, s, i)
            for e in s.els:
                self._el_pptx(sl, e)
            if s.notes:
                sl.notes_slide.notes_text_frame.text = s.notes
        _ensure_dir(os.path.dirname(path))
        prs.save(path)
        return path

    # ---------- pptx 内部 ----------
    def _chrome_pptx(self, sl, s, idx):
        bg = T["ink"] if s.dark else T["white"]
        sl.background.fill.solid()
        sl.background.fill.fore_color.rgb = RGBColor.from_string(bg)
        if s.dark:
            _rp(sl, 0, 0, SW, 0.10, fill=T["accent"])
        else:
            _rp(sl, 0, 0, SW, 0.085, fill=T["accent"])
        if s.kicker:
            _tb(sl, PAD, 0.40, 7.0, 0.26,
                [L(s.kicker, size=10.5, bold=True, color=T["accent"])])
        if s.title:
            _tb(sl, PAD, 0.64, SW - 2 * PAD, 0.56,
                [L(s.title, size=25, bold=True,
                   color=T["white"] if s.dark else T["ink"])])
        if s.sub:
            _tb(sl, PAD, 1.24, SW - 2 * PAD, 0.34,
                [L(s.sub, size=12, color=T["muted"])])
        if s.title and not s.dark:
            _ln(sl, PAD, 1.74, SW - 2 * PAD, T["line"], 1.0)
        if s.dark:
            _tb(sl, PAD, SH - 0.52, 8.0, 0.28,
                [L(self.footer_left, size=8.5, color="7E90AA")])
            _tb(sl, SW - PAD - 1.2, SH - 0.52, 1.2, 0.28,
                [L(f"{idx + 1:02d}", size=9, bold=True, color="7E90AA", align="r")])
        else:
            _tb(sl, PAD, SH - 0.52, 8.0, 0.28,
                [L(self.footer_left, size=8.5, color=T["muted"])])
            _tb(sl, SW - PAD - 1.2, SH - 0.52, 1.2, 0.28,
                [L(f"{idx + 1:02d}", size=9, bold=True, color=T["muted"], align="r")])

    def _el_pptx(self, sl, e):
        k = e["k"]
        if k == "rect":
            sh = _rp(sl, e["x"], e["y"], e["w"], e["h"], fill=e["fill"],
                     line=e["line"], lw=e["lw"], radius=e["radius"])
            if e.get("dash"):
                sh.line.dash_style = 4  # dash
        elif k in ("hline", "vline"):
            if k == "hline":
                _ln(sl, e["x"], e["y"], e["w"], e["color"], e["lw"])
            else:
                _lnv(sl, e["x"], e["y"], e["h"], e["color"], e["lw"])
        elif k == "text":
            _tb(sl, e["x"], e["y"], e["w"], e["h"], e["lines"],
                valign=e["valign"], pad=e.get("pad", 0.0))
        elif k == "chevron":
            sh = sl.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(e["x"]), Inches(e["y"]),
                                     Inches(e["w"]), Inches(e["h"]))
            sh.fill.solid()
            sh.fill.fore_color.rgb = RGBColor.from_string(e["fill"])
            sh.line.fill.background()
            sh.shadow.inherit = False
            tf = sh.text_frame
            tf.word_wrap = True
            _margins(tf, 0.04, 0.02)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            _run(p, e["text"], e["size"], e["bold"], e["tcolor"])
        elif k == "arrow":
            sh = sl.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(e["x"]), Inches(e["y"]),
                                     Inches(e["w"]), Inches(e["h"]))
            sh.fill.solid()
            sh.fill.fore_color.rgb = RGBColor.from_string(e["fill"])
            sh.line.fill.background()
            sh.shadow.inherit = False
        elif k == "table":
            _table_pptx(sl, e)

    def render(self, outdir, px_per_in=110):
        _ensure_dir(outdir)
        from PIL import Image, ImageDraw, ImageFont

        reg = ImageFont.truetype("/tmp/fonts/NotoSansSC-Regular.otf", 10)
        bold = ImageFont.truetype("/tmp/fonts/NotoSansSC-Bold.otf", 10)
        med = ImageFont.truetype("/tmp/fonts/NotoSansSC-Medium.otf", 10)
        cache: dict = {}

        def font(size, weight="r"):
            key = (round(size, 1), weight)
            if key not in cache:
                path = {"r": "/tmp/fonts/NotoSansSC-Regular.otf",
                        "m": "/tmp/fonts/NotoSansSC-Medium.otf",
                        "b": "/tmp/fonts/NotoSansSC-Bold.otf"}[weight]
                cache[key] = ImageFont.truetype(path, max(6, int(round(size * px_per_in / 72))))
            return cache[key]

        W, H = int(SW * px_per_in), int(SH * px_per_in)
        problems = []
        paths = []
        for idx, s in enumerate(self.slides):
            img = Image.new("RGB", (W, H), "#" + (T["ink"] if s.dark else T["white"]))
            d = ImageDraw.Draw(img)
            d.rectangle([0, 0, W, int(0.085 * px_per_in)], fill="#" + T["accent"])
            if s.kicker:
                _ptxt(d, PAD, 0.40, 7.0, 0.26, [L(s.kicker, size=10.5, bold=True,
                                                     color=T["accent"])], font, px_per_in)
            if s.title:
                _ptxt(d, PAD, 0.64, SW - 2 * PAD, 0.56,
                      [L(s.title, size=25, bold=True,
                         color=T["white"] if s.dark else T["ink"])], font, px_per_in)
            if s.sub:
                _ptxt(d, PAD, 1.24, SW - 2 * PAD, 0.34,
                      [L(s.sub, size=12, color=T["muted"])], font, px_per_in)
            if s.title and not s.dark:
                d.line([int(PAD * px_per_in), int(1.74 * px_per_in),
                        int((SW - PAD) * px_per_in), int(1.74 * px_per_in)],
                       fill="#" + T["line"], width=1)
            fc = "7E90AA" if s.dark else T["muted"]
            _ptxt(d, PAD, SH - 0.52, 8.0, 0.28,
                  [L(self.footer_left, size=8.5, color=fc)], font, px_per_in)
            _ptxt(d, SW - PAD - 1.2, SH - 0.52, 1.2, 0.28,
                  [L(f"{idx + 1:02d}", size=9, bold=True, color=fc, align="r")],
                  font, px_per_in)
            for e in s.els:
                prob = _el_pil(d, e, font, px_per_in, W, H)
                if prob:
                    problems.append(f"slide {idx+1}: {prob}")
                # 越界检查：任何元素不得越过脚注区（6.90in）或右/下边界
                bottom = e.get("y", 0) + (e.get("h", 0) if e["k"] != "table"
                                          else e["row_h"] * len(e["rows"]))
                right = e.get("x", 0) + e.get("w", 0)
                if bottom > 6.90:
                    problems.append(f"slide {idx+1}: 越过脚注区 bottom={bottom:.2f} :: "
                                    f"{e['k']} @ {e.get('x', 0):.2f},{e.get('y', 0):.2f}")
                if right > SW - 0.30:
                    problems.append(f"slide {idx+1}: 越过右边距 right={right:.2f} :: {e['k']}")
            p = os.path.join(outdir, f"slide-{idx+1:02d}.png")
            img.save(p)
            paths.append(p)
        return paths, problems


# ─────────────────────────── pptx helpers ───────────────────────────
def _ensure_dir(p):
    if p:
        os.makedirs(p, exist_ok=True)


def _rp(sl, x, y, w, h, fill=None, line=None, lw=1.0, radius=0.10):
    shp = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = sl.shapes.add_shape(shp, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        try:
            sh.adjustments[0] = min(0.5, radius / max(0.01, min(w, h)) * 0.6)
        except Exception:
            pass
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = RGBColor.from_string(line)
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def _ln(sl, x, y, w, color, lw=1.0):
    """水平分隔线。

    注意：这里刻意不用 add_connector —— 连接线在水平方向时 cy=0，
    属于 zero-extent 形状。PowerPoint 容错，但 Keynote 会判定整个文件
    「file format is invalid」。改用极细矩形，兼容性最好。
    """
    h = max(0.012, lw / 72.0)
    sh = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y - h / 2),
                             Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(color)
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def _lnv(sl, x, y, h, color, lw=1.0):
    """垂直分隔线（同理：不用连接线）。"""
    w = max(0.012, lw / 72.0)
    sh = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x - w / 2), Inches(y),
                             Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(color)
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def _margins(tf, lr=0.0, tb=0.0):
    tf.margin_left = Pt(lr * 72)
    tf.margin_right = Pt(lr * 72)
    tf.margin_top = Pt(tb * 72)
    tf.margin_bottom = Pt(tb * 72)


def _run(p, text, size, bold, color, ):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    r.font.name = FONT_LAT
    # 东亚字体
    rPr = r._r.get_or_add_rPr()
    from pptx.oxml.ns import qn
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", FONT_CJK)
    return r


def _tb(sl, x, y, w, h, lines, valign="t", pad=0.0):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    _margins(tf, pad, 0.01)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE,
                          "b": MSO_ANCHOR.BOTTOM}[valign]
    first = True
    for ln in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER,
                       "r": PP_ALIGN.RIGHT}[ln.get("align", "l")]
        if ln.get("space"):
            p.space_before = Pt(ln["space"])
        if ln.get("line"):
            p.line_spacing = ln["line"]
        txt = ln["t"]
        if ln.get("bullet"):
            txt = f"{ln['bullet']}  {txt}"
        _run(p, txt, ln["size"], ln["bold"], ln["color"])
    return tb


def _table_pptx(sl, e):
    rows, cols = len(e["rows"]), len(e["rows"][0])
    gt = sl.shapes.add_table(rows, cols, Inches(e["x"]), Inches(e["y"]),
                             Inches(e["w"]), Inches(e["row_h"] * rows)).table
    gt.first_row = e["header"]
    total = sum(e["col_w"])
    for j, cw in enumerate(e["col_w"]):
        gt.columns[j].width = Emu(int(Inches(e["w"]) * cw / total))
    for i, row in enumerate(e["rows"]):
        gt.rows[i].height = Inches(e["row_h"])
        for j, cell in enumerate(row):
            c = gt.cell(i, j)
            c.margin_left = Pt(5)
            c.margin_right = Pt(4)
            c.margin_top = Pt(1)
            c.margin_bottom = Pt(1)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            is_head = e["header"] and i == 0
            txt = cell if isinstance(cell, str) else cell.get("t", "")
            col = (T["white"] if is_head
                   else (cell.get("c", T["body"]) if isinstance(cell, dict) else T["body"]))
            bold = is_head or (isinstance(cell, dict) and cell.get("b", False)) or j == 0
            _run(p, txt, e["size"], bold, col)
            c.fill.solid()
            if is_head:
                c.fill.fore_color.rgb = RGBColor.from_string(e["head_fill"])
            elif e["zebra"] and i % 2 == 0:
                c.fill.fore_color.rgb = RGBColor.from_string(T["soft"])
            else:
                c.fill.fore_color.rgb = RGBColor.from_string(T["white"])
    return gt


# ─────────────────────────── PIL helpers ───────────────────────────
CJK_RANGES = ((0x2E80, 0x9FFF), (0xF900, 0xFAFF), (0xFF00, 0xFFEF),
              (0x3000, 0x303F))


def _is_cjk(ch):
    o = ord(ch)
    return any(a <= o <= b for a, b in CJK_RANGES)


def _wrap(text, font, max_px):
    """按像素宽度折行：CJK 逐字断，拉丁按空格断。"""
    out, cur = [], ""
    tokens = []
    buf = ""
    for ch in text:
        if ch in " \t":
            if buf:
                tokens.append(buf)
                buf = ""
            tokens.append(ch)
        elif _is_cjk(ch):
            if buf:
                tokens.append(buf)
                buf = ""
            tokens.append(ch)
        else:
            buf += ch
    if buf:
        tokens.append(buf)
    for tk in tokens:
        trial = cur + tk
        if font.getlength(trial.strip()) <= max_px or not cur.strip():
            cur = trial
        else:
            out.append(cur.rstrip())
            cur = "" if tk == " " else tk
    if cur.strip():
        out.append(cur.rstrip())
    return out or [""]


def _ptxt(d, x, y, w, h, lines, font, ppi, valign=None):
    px = lambda v: int(v * ppi)
    cy = y
    total = 0.0
    for ln in lines:
        f = font(ln["size"], "b" if ln["bold"] else "r")
        wrapped = _wrap(ln["t"], f, max(8, px(w) - 2))
        lh = ln["size"] * 1.42 / 72
        if ln.get("space"):
            cy += ln["space"] / 72
            total += ln["space"] / 72
        for wl in wrapped:
            if valign == "m":
                pass
            ax = x
            if ln.get("align") == "c":
                ax = x + (w - f.getlength(wl) / ppi) / 2
            elif ln.get("align") == "r":
                ax = x + w - f.getlength(wl) / ppi
            d.text((px(ax), px(cy)), wl, font=f, fill="#" + ln["color"])
            cy += lh
            total += lh
    return total


def _el_pil(d, e, font, ppi, W, H):
    px = lambda v: int(v * ppi)
    k = e["k"]
    if k == "rect":
        if e.get("fill"):
            d.rounded_rectangle([px(e["x"]), px(e["y"]), px(e["x"] + e["w"]),
                                 px(e["y"] + e["h"])],
                                radius=px(e["radius"]) if e["radius"] else 0,
                                fill="#" + e["fill"])
        if e.get("line"):
            d.rounded_rectangle([px(e["x"]), px(e["y"]), px(e["x"] + e["w"]),
                                 px(e["y"] + e["h"])],
                                radius=px(e["radius"]) if e["radius"] else 0,
                                outline="#" + e["line"], width=max(1, int(e["lw"])))
    elif k == "hline":
        d.line([px(e["x"]), px(e["y"]), px(e["x"] + e["w"]), px(e["y"])],
               fill="#" + e["color"], width=max(1, int(e["lw"])))
    elif k == "vline":
        d.line([px(e["x"]), px(e["y"]), px(e["x"]), px(e["y"] + e["h"])],
               fill="#" + e["color"], width=max(1, int(e["lw"])))
    elif k == "text":
        used = _ptxt(d, e["x"], e["y"], e["w"], e["h"], e["lines"], font, ppi,
                     valign=e.get("valign"))
        if used > e["h"] + 0.02:
            return f"文本溢出 {used - e['h']:.2f}in :: {e['lines'][0]['t'][:32]}"
    elif k == "chevron":
        x0, y0, x1, y1 = px(e["x"]), px(e["y"]), px(e["x"] + e["w"]), px(e["y"] + e["h"])
        notch = (y1 - y0) * 0.42
        pts = [(x0, y0), (x1 - notch, y0), (x1, (y0 + y1) / 2), (x1 - notch, y1),
               (x0, y1), (x0 + notch, (y0 + y1) / 2)]
        d.polygon(pts, fill="#" + e["fill"])
        f = font(e["size"], "b" if e["bold"] else "r")
        tw = f.getlength(e["text"])
        d.text(((x0 + notch + x1 - notch) / 2 - tw / 2, (y0 + y1) / 2 - e["size"] * 0.72),
               e["text"], font=f, fill="#" + e["tcolor"])
    elif k == "arrow":
        x0, y0, x1, y1 = px(e["x"]), px(e["y"]), px(e["x"] + e["w"]), px(e["y"] + e["h"])
        mid = (y0 + y1) / 2
        head = min((x1 - x0) * 0.5, (y1 - y0) * 1.2)
        d.rectangle([x0, mid - (y1 - y0) * 0.16, x1 - head, mid + (y1 - y0) * 0.16],
                    fill="#" + e["fill"])
        d.polygon([(x1 - head, y0), (x1, mid), (x1 - head, y1)], fill="#" + e["fill"])
    elif k == "table":
        _table_pil(d, e, font, ppi)
    return None


def _table_pil(d, e, font, ppi):
    px = lambda v: int(v * ppi)
    total = sum(e["col_w"])
    widths = [px(e["w"]) * cw / total for cw in e["col_w"]]
    y = e["y"]
    for i, row in enumerate(e["rows"]):
        is_head = e["header"] and i == 0
        bg = e["head_fill"] if is_head else (T["soft"] if e["zebra"] and i % 2 == 0
                                             else T["white"])
        d.rectangle([px(e["x"]), px(y), px(e["x"] + e["w"]), px(y + e["row_h"])],
                    fill="#" + bg)
        x = e["x"]
        for j, cell in enumerate(row):
            txt = cell if isinstance(cell, str) else cell.get("t", "")
            col = (T["white"] if is_head
                   else (cell.get("c", T["body"]) if isinstance(cell, dict) else T["body"]))
            bold = is_head or (isinstance(cell, dict) and cell.get("b")) or j == 0
            f = font(e["size"], "b" if bold else "r")
            wrapped = _wrap(txt, f, max(10, widths[j] - px(0.14)))
            lh = e["size"] * 1.30 / 72
            ty = y + (e["row_h"] - lh * len(wrapped)) / 2
            for wl in wrapped[:2]:
                d.text((px(x + 0.07), px(ty)), wl, font=f, fill="#" + col)
                ty += lh
            x += widths[j] / ppi
        y += e["row_h"]
        d.line([px(e["x"]), px(y), px(e["x"] + e["w"]), px(y)], fill="#" + T["line"], width=1)


def contact_sheet(paths, out, cols=4, title=None):
    from PIL import Image, ImageDraw, ImageFont
    ims = [Image.open(p) for p in paths]
    w, h = ims[0].size
    sc = 0.42
    tw, th = int(w * sc), int(h * sc)
    rows = (len(ims) + cols - 1) // cols
    pad, gap = 16, 12
    headh = 44 if title else 0
    sheet = Image.new("RGB", (pad * 2 + cols * tw + (cols - 1) * gap,
                              headh + pad * 2 + rows * th + (rows - 1) * gap), "#E9EDF3")
    d = ImageDraw.Draw(sheet)
    if title:
        f = ImageFont.truetype("/tmp/fonts/NotoSansSC-Bold.otf", 24)
        d.text((pad, 12), title, font=f, fill="#0F1B2D")
    for i, im in enumerate(ims):
        r, c = divmod(i, cols)
        x = pad + c * (tw + gap)
        y = headh + pad + r * (th + gap)
        sheet.paste(im.resize((tw, th), Image.LANCZOS), (x, y))
        d.rectangle([x - 1, y - 1, x + tw, y + th], outline="#C6CFDC", width=1)
    sheet.save(out)
    return out


# ─────────────────── 兼容性体检（Keynote 尤其挑剔） ───────────────────
def audit_pptx(path):
    """检查 import 兼容性风险，返回问题列表。

    Keynote 对以下情况会直接报 "file format is invalid"：
      · zero-extent 形状（cx=0 或 cy=0）—— 最常见
      · p:cxnSp 连接线
      · 空文本 run
      · rPr 子元素顺序错误
    """
    import zipfile, re
    problems = []
    z = zipfile.ZipFile(path)
    slides = [n for n in z.namelist() if n.startswith("ppt/slides/slide")]
    zero = conn = emptyt = 0
    for n in slides:
        x = z.read(n).decode("utf-8")
        for m in re.finditer(r'<a:ext cx="(-?\d+)" cy="(-?\d+)"/>', x):
            if int(m.group(1)) <= 0 or int(m.group(2)) <= 0:
                zero += 1
        conn += x.count("<p:cxnSp>")
        emptyt += len(re.findall(r'<a:t>\s*</a:t>', x))
    if zero:
        problems.append(f"zero-extent 形状 {zero} 个（Keynote 会拒绝）")
    if conn:
        problems.append(f"连接线 p:cxnSp {conn} 个（建议改为矩形）")
    if emptyt:
        problems.append(f"空文本 run {emptyt} 个")
    return problems
